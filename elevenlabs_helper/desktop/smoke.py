"""Offline checks run against frozen release binaries on their native OS."""

from pathlib import Path
from importlib import import_module
import logging
import os
import sys
import wave
from tempfile import TemporaryDirectory


def main(result_path: str) -> int:
    from docx import Document
    from PySide6.QtWidgets import QApplication

    from ..engine.service import Engine
    from ..engine.media.inspect import inspect
    from .bridge import EngineBridge
    from .widgets.key_dialog import KeyDialog
    from .windows.main_window import MainWindow

    with TemporaryDirectory() as temporary:
        os.environ["XDG_DATA_HOME"] = temporary
        os.environ["APPDATA"] = temporary
        os.environ["ELEVENLABS_API_KEY"] = "offline-smoke-test"
        backend = {
            "win32": "keyring.backends.Windows.WinVaultKeyring",
            "linux": "keyring.backends.SecretService.Keyring",
            "darwin": "keyring.backends.macOS.Keyring",
        }[sys.platform]
        module, _, name = backend.rpartition(".")
        assert getattr(import_module(module), name)
        app = QApplication([])
        bridge = EngineBridge()
        engine = Engine(on_update=bridge.push)
        window = MainWindow(engine, bridge)
        dialog = KeyDialog(window)
        window.show()
        app.processEvents()
        assert not window.grab().isNull()
        assert dialog.field.text() == "offline-smoke-test"
        document = Document()
        document.add_paragraph("Offline package check")
        output = Path(temporary) / "check.docx"
        document.save(output)
        assert Document(output).paragraphs[0].text == "Offline package check"
        wav = Path(temporary) / "check.WAV"
        with wave.open(str(wav), "wb") as audio:
            audio.setnchannels(1)
            audio.setsampwidth(2)
            audio.setframerate(8000)
            audio.writeframes(b"\x00\x00" * 8000)
        assert inspect(wav).duration_secs == 1.0
        assert window.drop._is_media(str(wav))
        window.close()
        logging.shutdown()
        Path(result_path).write_text("GUI, credential backend, SQLite, DOCX, and WAV: OK\n")
    return 0
